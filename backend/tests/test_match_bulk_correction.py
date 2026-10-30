"""Bulk coefficient and correction import tests."""
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api import match_api
from app.api.match_api import router
from app.core.auth_deps import get_current_user
from app.models.database import Base, get_db
from app.models.schemas import (
    Category,
    CleanJobRecord,
    MatchResult,
    ModelRecord,
    RawDataRecord,
    UploadFileRecord,
)


class DummyUser:
    is_admin = 1
    permissions = []
    category_permissions = []


@pytest.fixture
def client(monkeypatch):
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)

    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_db] = lambda: (lambda: None)()
    # Provide a real session per request.
    def override_db():
        session = Session()
        try:
            yield session
        finally:
            session.close()
    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = lambda: DummyUser()
    monkeypatch.setattr(match_api, "audit_price", lambda *args, **kwargs: {"audited": 0})
    test_client = TestClient(app)
    test_client.Session = Session
    return test_client


def _seed(client, price=100):
    with client.Session() as db:
        db.add(Category(code="headphone", name="耳机"))
        db.add(ModelRecord(brand_code="BOSE", model_code="SB900", category_code="headphone", brand_name="Bose", model_name="Soundbar"))
        upload = UploadFileRecord(filename="test.xlsx", platform="jd", row_count=1, status="done")
        db.add(upload)
        db.flush()
        job = CleanJobRecord(file_ids=[upload.id], status="reviewing", category_code="headphone")
        db.add(job)
        db.flush()
        raw = RawDataRecord(
            file_id=upload.id, platform="jd", month=202609, item_id="item-1",
            item_url="https://item.jd.com/item-1.html",
            item_name="Original name", price=price, sales_qty=10, sales_amount=1000,
        )
        db.add(raw)
        db.flush()
        model = db.query(ModelRecord).filter_by(model_code="SB900").one()
        match = MatchResult(
            clean_job_id=job.id, raw_data_id=raw.id, model_id=model.id,
            match_status="matched", matched_by="auto", match_source="s1",
            adjusted_price=110,
        )
        db.add(match)
        db.commit()
        return job.id, match.id


def test_bulk_sales_coefficient(client):
    job_id, match_id = _seed(client)
    res = client.post("/api/match/reviewed/bulk-coefficient", json={
        "clean_job_id": job_id, "tab": "all", "target": "sales", "coefficient": 1.1,
    })
    assert res.status_code == 200, res.text
    assert res.json()["updated_count"] == 1
    with client.Session() as db:
        assert float(db.query(MatchResult).filter_by(id=match_id).one().sales_coefficient) == 1.1


def test_bulk_price_coefficient(client):
    job_id, match_id = _seed(client, price=100)
    res = client.post("/api/match/reviewed/bulk-coefficient", json={
        "clean_job_id": job_id, "tab": "all", "target": "price", "coefficient": 0.9,
    })
    assert res.status_code == 200, res.text
    with client.Session() as db:
        assert float(db.query(MatchResult).filter_by(id=match_id).one().adjusted_price) == 90.0


def test_bulk_sales_coefficient_respects_excluded_ids(client):
    job_id, first_match_id = _seed(client)
    with client.Session() as db:
        job = db.query(CleanJobRecord).filter_by(id=job_id).one()
        upload = db.query(UploadFileRecord).first()
        model = db.query(ModelRecord).filter_by(model_code="SB900").one()
        raw = RawDataRecord(
            file_id=upload.id, platform="jd", month=202609, item_id="item-2",
            item_name="Second name", price=50, sales_qty=5, sales_amount=250,
        )
        db.add(raw)
        db.flush()
        second = MatchResult(
            clean_job_id=job.id, raw_data_id=raw.id, model_id=model.id,
            match_status="matched", matched_by="auto", match_source="s1",
        )
        db.add(second)
        db.commit()
        second_match_id = second.id

    res = client.post("/api/match/reviewed/bulk-coefficient", json={
        "clean_job_id": job_id, "tab": "all", "target": "sales",
        "coefficient": 1.2, "exclude_ids": [first_match_id],
    })
    assert res.status_code == 200, res.text
    assert res.json()["updated_count"] == 1
    with client.Session() as db:
        assert db.query(MatchResult).filter_by(id=first_match_id).one().sales_coefficient is None
        assert float(db.query(MatchResult).filter_by(id=second_match_id).one().sales_coefficient) == 1.2


def _xlsx_bytes(rows):
    import io
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def test_export_reviewed_matches_includes_category(client):
    import io
    from openpyxl import load_workbook

    job_id, _match_id = _seed(client)
    res = client.get("/api/match/reviewed/export", params={"clean_job_id": job_id})
    assert res.status_code == 200, res.text
    workbook = load_workbook(io.BytesIO(res.content))
    sheet = workbook["匹配结果"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert "品类" in headers
    assert "调整后价格" in headers
    values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    row = dict(zip(headers, values))
    assert row["品类"] == "耳机"
    assert row["商品ID"] == "item-1"


def test_bulk_update_by_url_recalculates_coefficients(client):
    job_id, match_id = _seed(client)
    content = _xlsx_bytes([
        ["商品链接", "销量", "价格"],
        ["https://item.jd.com/item-1.html", 25, 88.5],
    ])
    res = client.post(
        "/api/match/reviewed/bulk-update-by-url",
        data={"clean_job_id": str(job_id)},
        files={"file": ("bulk.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["updated_count"] == 1
    assert body["skipped_rows"] == 0
    with client.Session() as db:
        match = db.query(MatchResult).filter_by(id=match_id).one()
        assert float(match.sales_coefficient) == 2.5
        assert float(match.adjusted_price) == 88.5


def test_bulk_update_by_url_reports_missing_url(client):
    job_id, _match_id = _seed(client)
    content = _xlsx_bytes([
        ["商品链接", "销量"],
        ["https://item.jd.com/not-found.html", 1],
    ])
    res = client.post(
        "/api/match/reviewed/bulk-update-by-url",
        data={"clean_job_id": str(job_id)},
        files={"file": ("bulk.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["updated_count"] == 0
    assert body["skipped_rows"] == 1
    assert body["errors"]


def test_bulk_update_template_download(client):
    res = client.get("/api/match/reviewed/bulk-update-template")
    assert res.status_code == 200
    assert res.content[:2] == b"PK"
