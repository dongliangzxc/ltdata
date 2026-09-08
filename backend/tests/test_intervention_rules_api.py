from fastapi import FastAPI
from fastapi.testclient import TestClient
from types import SimpleNamespace

from app.api.rules_api import router
from app.core.auth_deps import get_current_user
from app.models.database import get_db
from app.models.schemas import (
    Category,
    CleanJobRecord,
    FilteredItem,
    InterventionRule,
    RawDataRecord,
    UploadFileRecord,
)


def _make_client(db):
    app = FastAPI()
    app.include_router(router)

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(
        is_admin=1, category_permissions=None, username="tester",
    )
    return TestClient(app)


def test_create_and_list_intervention_rule(db):
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    db.commit()

    response = client.post("/api/rules/intervention-rules", json={
        "name": "海信低价配件过滤",
        "category_code": "projector",
        "action": "filter",
        "priority": 10,
        "conditions": {
            "brand_in": ["海信"],
            "item_name_not_contains_any": ["激光电视"],
            "reference_price": {"op": "lt", "value": 500},
        },
    })

    assert response.status_code == 201
    created = response.json()
    assert created["id"] > 0
    assert created["name"] == "海信低价配件过滤"
    assert created["category_code"] == "projector"
    assert created["action"] == "filter"
    assert created["priority"] == 10
    assert created["is_active"] == 1
    assert created["conditions"]["brand_in"] == ["海信"]

    listed = client.get("/api/rules/intervention-rules", params={"category_code": "projector"}).json()
    assert [row["id"] for row in listed] == [created["id"]]
    assert listed[0]["summary"] == "入库品牌 in [海信] 且 商品名称不包含 [激光电视] 且 参考价格 < 500"


def test_intervention_rule_requires_category_and_valid_action(db):
    client = _make_client(db)

    missing_category = client.post("/api/rules/intervention-rules", json={
        "name": "无品类规则",
        "category_code": "",
        "action": "filter",
        "priority": 10,
        "conditions": {"item_name_contains_any": ["配件"]},
    })
    assert missing_category.status_code == 400
    assert missing_category.json()["detail"] == "category_code 不能为空"

    db.add(Category(code="projector", name="投影"))
    db.commit()

    invalid_action = client.post("/api/rules/intervention-rules", json={
        "name": "错误动作",
        "category_code": "projector",
        "action": "delete",
        "priority": 10,
        "conditions": {"item_name_contains_any": ["配件"]},
    })
    assert invalid_action.status_code == 400
    assert invalid_action.json()["detail"] == "action 必须是 filter 或 allow"


def test_intervention_rule_validates_price_condition(db):
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    db.commit()

    response = client.post("/api/rules/intervention-rules", json={
        "name": "错误价格规则",
        "category_code": "projector",
        "action": "filter",
        "priority": 10,
        "conditions": {"reference_price": {"op": "between", "value": 500}},
    })

    assert response.status_code == 400
    assert response.json()["detail"] == "reference_price between 必须包含 min 和 max"


def test_intervention_rule_rejects_unknown_condition_key(db):
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    db.commit()

    response = client.post("/api/rules/intervention-rules", json={
        "name": "未知条件规则",
        "category_code": "projector",
        "action": "filter",
        "priority": 10,
        "conditions": {"unknown_key": ["x"]},
    })

    assert response.status_code == 400
    assert response.json()["detail"] == "不支持的干预条件: unknown_key"


def test_intervention_rule_rejects_invalid_numeric_price_value(db):
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    db.commit()

    response = client.post("/api/rules/intervention-rules", json={
        "name": "非法价格规则",
        "category_code": "projector",
        "action": "filter",
        "priority": 10,
        "conditions": {"reference_price": {"op": "lt", "value": "abc"}},
    })

    assert response.status_code == 400
    assert response.json()["detail"] == "reference_price 数值必须是有效数字"


def test_intervention_rule_rejects_non_finite_price_value(db):
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    db.commit()

    response = client.post("/api/rules/intervention-rules", json={
        "name": "非有限价格规则",
        "category_code": "projector",
        "action": "filter",
        "priority": 10,
        "conditions": {"reference_price": {"op": "lt", "value": "nan"}},
    })

    assert response.status_code == 400
    assert response.json()["detail"] == "reference_price 数值必须是有效数字"


def test_intervention_rule_rejects_between_min_greater_than_max(db):
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    db.commit()

    response = client.post("/api/rules/intervention-rules", json={
        "name": "价格区间错误规则",
        "category_code": "projector",
        "action": "filter",
        "priority": 10,
        "conditions": {"reference_price": {"op": "between", "min": 1000, "max": 500}},
    })

    assert response.status_code == 400
    assert response.json()["detail"] == "reference_price between 最低价不能大于最高价"


def test_intervention_rule_patch_rejects_invalid_is_active(db):
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    db.add(InterventionRule(
        name="旧规则",
        category_code="projector",
        action="filter",
        priority=50,
        conditions={"item_name_contains_any": ["配件"]},
    ))
    db.commit()
    rule = db.query(InterventionRule).first()

    response = client.patch(f"/api/rules/intervention-rules/{rule.id}", json={"is_active": 2})

    assert response.status_code == 400
    assert response.json()["detail"] == "is_active 必须是 0 或 1"


def test_update_toggle_and_delete_intervention_rule(db):
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    db.add(InterventionRule(
        name="旧规则",
        category_code="projector",
        action="filter",
        priority=50,
        conditions={"item_name_contains_any": ["配件"]},
    ))
    db.commit()
    rule = db.query(InterventionRule).first()

    updated = client.patch(f"/api/rules/intervention-rules/{rule.id}", json={
        "name": "新规则",
        "action": "allow",
        "priority": 5,
        "is_active": 0,
    })
    assert updated.status_code == 200
    assert updated.json()["name"] == "新规则"
    assert updated.json()["action"] == "allow"
    assert updated.json()["priority"] == 5
    assert updated.json()["is_active"] == 0

    deleted = client.delete(f"/api/rules/intervention-rules/{rule.id}")
    assert deleted.status_code == 204
    assert db.query(InterventionRule).count() == 0


def test_filtered_items_include_intervention_rule_reason(db):
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    rule = InterventionRule(
        name="配件过滤",
        category_code="projector",
        action="filter",
        priority=10,
        conditions={"item_name_contains_any": ["配件"]},
    )
    upload_file = UploadFileRecord(filename="test.xlsx", platform="JD", row_count=1)
    db.add_all([rule, upload_file])
    db.flush()

    raw = RawDataRecord(
        file_id=upload_file.id,
        platform="JD",
        month=202605,
        item_id="sku-1",
        item_name="投影仪配件",
        brand_raw="测试品牌",
        shop_name="测试店铺",
    )
    clean_job = CleanJobRecord(file_ids=[upload_file.id], rules={}, status="done", row_in=1, row_out=0, row_filtered=1)
    db.add_all([raw, clean_job])
    db.flush()

    db.add(FilteredItem(
        raw_data_id=raw.id,
        clean_job_id=clean_job.id,
        matched_keyword="配件",
        intervention_rule_id=rule.id,
        intervention_rule_name=rule.name,
        matched_reason="命中规则「配件过滤」：商品名称包含 [配件]",
    ))
    db.commit()

    response = client.get("/api/rules/filtered-items")

    assert response.status_code == 200
    first_item = response.json()["items"][0]
    assert first_item["intervention_rule_id"] == rule.id
    assert first_item["intervention_rule_name"] == "配件过滤"
    assert first_item["matched_reason"] == "命中规则「配件过滤」：商品名称包含 [配件]"


def _seed_filtered_row(db, *, clean_job_id, upload_id, item_name, brand_raw, rule):
    raw = RawDataRecord(
        file_id=upload_id,
        platform="JD",
        month=202605,
        item_id=f"sku-{item_name}",
        item_name=item_name,
        brand_raw=brand_raw,
        shop_name="店铺",
    )
    db.add(raw)
    db.flush()
    fi = FilteredItem(
        raw_data_id=raw.id,
        clean_job_id=clean_job_id,
        matched_keyword="配件",
        intervention_rule_id=rule.id,
        intervention_rule_name=rule.name,
        matched_reason=f"命中规则「{rule.name}」",
    )
    db.add(fi)
    db.flush()
    return fi


def test_filtered_items_search_by_brand_raw(db):
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    rule = InterventionRule(
        name="配件过滤",
        category_code="projector",
        action="filter",
        priority=10,
        conditions={"item_name_contains_any": ["配件"]},
    )
    upload_file = UploadFileRecord(filename="test.xlsx", platform="JD", row_count=2)
    db.add_all([rule, upload_file])
    db.flush()
    clean_job = CleanJobRecord(file_ids=[upload_file.id], rules={}, status="done", row_in=2, row_out=0, row_filtered=2)
    db.add(clean_job)
    db.flush()

    hit = _seed_filtered_row(
        db, clean_job_id=clean_job.id, upload_id=upload_file.id,
        item_name="投影仪配件A", brand_raw="大疆DJI", rule=rule,
    )
    _seed_filtered_row(
        db, clean_job_id=clean_job.id, upload_id=upload_file.id,
        item_name="投影仪配件B", brand_raw="SONY", rule=rule,
    )
    db.commit()

    resp = client.get(
        "/api/rules/filtered-items",
        params={"clean_job_id": clean_job.id, "search_by": "brand_raw", "keyword": "大疆"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 1
    assert body["items"][0]["id"] == hit.id


def test_filtered_items_search_by_item_name(db):
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    rule = InterventionRule(
        name="配件过滤",
        category_code="projector",
        action="filter",
        priority=10,
        conditions={"item_name_contains_any": ["配件"]},
    )
    upload_file = UploadFileRecord(filename="t.xlsx", platform="JD", row_count=2)
    db.add_all([rule, upload_file])
    db.flush()
    clean_job = CleanJobRecord(file_ids=[upload_file.id], rules={}, status="done", row_in=2, row_out=0, row_filtered=2)
    db.add(clean_job)
    db.flush()

    hit = _seed_filtered_row(
        db, clean_job_id=clean_job.id, upload_id=upload_file.id,
        item_name="Mavic 3 配件", brand_raw="DJI", rule=rule,
    )
    _seed_filtered_row(
        db, clean_job_id=clean_job.id, upload_id=upload_file.id,
        item_name="Osmo 配件", brand_raw="DJI", rule=rule,
    )
    db.commit()

    resp = client.get(
        "/api/rules/filtered-items",
        params={"clean_job_id": clean_job.id, "search_by": "item_name", "keyword": "Mavic"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 1
    assert body["items"][0]["id"] == hit.id


def test_filtered_items_keyword_no_longer_matches_matched_keyword_or_rule_name(db):
    """本次改造后，keyword 只按 search_by 指定的字段搜；不再对 matched_keyword / intervention_rule_name / matched_reason 做 OR 匹配。"""
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    rule = InterventionRule(
        name="配件过滤",
        category_code="projector",
        action="filter",
        priority=10,
        conditions={"item_name_contains_any": ["配件"]},
    )
    upload_file = UploadFileRecord(filename="t.xlsx", platform="JD", row_count=1)
    db.add_all([rule, upload_file])
    db.flush()
    clean_job = CleanJobRecord(file_ids=[upload_file.id], rules={}, status="done", row_in=1, row_out=0, row_filtered=1)
    db.add(clean_job)
    db.flush()

    _seed_filtered_row(
        db, clean_job_id=clean_job.id, upload_id=upload_file.id,
        item_name="Mavic 3", brand_raw="DJI", rule=rule,
    )  # matched_keyword="配件"，intervention_rule_name="配件过滤"，matched_reason 含"配件过滤"
    db.commit()

    # 默认按 item_name 搜"配件过滤"，item_name 里没有该子串，应返回 0 条
    resp = client.get(
        "/api/rules/filtered-items",
        params={"clean_job_id": clean_job.id, "keyword": "配件过滤"},
    )
    assert resp.status_code == 200
    assert resp.json()["total"] == 0


# ── 店铺名称条件 ─────────────────────────────────────────────

def test_intervention_rule_accepts_shop_name_condition(db):
    client = _make_client(db)
    db.add(Category(code="projector", name="投影"))
    db.commit()

    response = client.post("/api/rules/intervention-rules", json={
        "name": "指定店铺过滤",
        "category_code": "projector",
        "action": "filter",
        "priority": 1,
        "conditions": {"shop_name_in": ["京东某专营店", "天猫某旗舰店"]},
    })

    assert response.status_code == 201
    created = response.json()
    assert created["conditions"]["shop_name_in"] == ["京东某专营店", "天猫某旗舰店"]
    assert created["summary"] == "店铺名称 in [京东某专营店, 天猫某旗舰店]"


# ── 干扰链接库 ───────────────────────────────────────────────

def test_import_interference_links_from_excel(db):
    import io
    import openpyxl
    from app.models.schemas import InterferenceLink

    client = _make_client(db)
    db.add(Category(code="tv", name="电视"))
    db.commit()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "链接"
    ws.append(["链接", "品类", "备注"])
    ws.append(["https://item.jd.com/1001.html", "tv", "历史干扰"])
    ws.append(["https://item.jd.com/1002.html", "tv"])
    ws.append(["https://item.jd.com/1001.html", "tv"])  # 重复
    buf = io.BytesIO()
    wb.save(buf)

    resp = client.post(
        "/api/rules/interference-links/import",
        files={"file": ("links.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["imported"] == 2
    assert data["skipped"] == 1

    rows = db.query(InterferenceLink).order_by(InterferenceLink.id).all()
    assert [r.url for r in rows] == [
        "https://item.jd.com/1001.html",
        "https://item.jd.com/1002.html",
    ]
    assert rows[0].remark == "历史干扰"


def test_import_interference_links_requires_category_column(db):
    import io
    import openpyxl

    client = _make_client(db)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["链接", "备注"])
    ws.append(["https://item.jd.com/3000.html", "无品类列"])
    buf = io.BytesIO()
    wb.save(buf)

    resp = client.post(
        "/api/rules/interference-links/import",
        files={"file": ("links.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 400
    assert "品类" in resp.json()["detail"]


def test_import_interference_links_requires_url_column(db):
    import io
    import openpyxl

    client = _make_client(db)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["其他列"])
    ws.append(["x"])
    buf = io.BytesIO()
    wb.save(buf)

    resp = client.post(
        "/api/rules/interference-links/import",
        files={"file": ("links.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 400
    assert "链接" in resp.json()["detail"]


def test_list_and_delete_interference_links(db):
    from app.models.schemas import InterferenceLink

    client = _make_client(db)
    db.add(Category(code="tv", name="电视"))
    db.add(InterferenceLink(url="https://item.jd.com/2001.html", category_code="tv", remark="备注A"))
    db.add(InterferenceLink(url="https://item.jd.com/2002.html", category_code="tv"))
    db.commit()

    resp = client.get("/api/rules/interference-links")
    assert resp.status_code == 200
    data = resp.json()
    assert data["total"] == 2
    assert len(data["items"]) == 2

    link_id = data["items"][0]["id"]
    assert client.delete(f"/api/rules/interference-links/{link_id}").status_code == 204
    remaining = db.query(InterferenceLink).count()
    assert remaining == 1


def test_import_interference_links_with_category_column(db):
    import io
    import openpyxl
    from app.models.schemas import InterferenceLink

    client = _make_client(db)
    db.add(Category(code="tv", name="电视"))
    db.commit()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "链接"
    ws.append(["链接", "品类", "备注"])
    ws.append(["https://item.jd.com/3001.html", "tv", "电视专用"])
    ws.append(["https://item.jd.com/3002.html", "电视", "中文品类名"])  # 名称解析
    ws.append(["https://item.jd.com/3003.html", "", "空品类报错"])
    ws.append(["https://item.jd.com/3004.html", "不存在的品类", "报错"])
    buf = io.BytesIO()
    wb.save(buf)

    resp = client.post(
        "/api/rules/interference-links/import",
        files={"file": ("links.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["imported"] == 2
    assert data["skipped"] == 2
    assert len(data["errors"]) == 2
    assert any("品类不能为空" in e for e in data["errors"])
    assert any("无法识别品类" in e for e in data["errors"])

    rows = db.query(InterferenceLink).order_by(InterferenceLink.url).all()
    by_url = {r.url: r for r in rows}
    assert by_url["https://item.jd.com/3001.html"].category_code == "tv"
    assert by_url["https://item.jd.com/3002.html"].category_code == "tv"
    assert "https://item.jd.com/3003.html" not in by_url
    assert "https://item.jd.com/3004.html" not in by_url


def test_list_interference_links_includes_category_name(db):
    from app.models.schemas import InterferenceLink

    client = _make_client(db)
    db.add(Category(code="tv", name="电视"))
    db.add(Category(code="monitor", name="显示器"))
    db.add(InterferenceLink(url="https://item.jd.com/4001.html", category_code="tv"))
    db.add(InterferenceLink(url="https://item.jd.com/4002.html", category_code="monitor"))
    db.commit()

    data = client.get("/api/rules/interference-links").json()
    by_url = {i["url"]: i for i in data["items"]}
    assert by_url["https://item.jd.com/4001.html"]["category_name"] == "电视"
    assert by_url["https://item.jd.com/4002.html"]["category_name"] == "显示器"

    filtered = client.get("/api/rules/interference-links", params={"category_code": "tv"}).json()
    assert filtered["total"] == 1
    assert filtered["items"][0]["url"] == "https://item.jd.com/4001.html"
